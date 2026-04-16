import logging
from typing import Dict, Any, List
from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

from schemas.output_schema import SectionData, ParsedDocument
import uuid

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


class DocxParser:
    def __init__(self):
        self.heading_styles = [
            "Heading 1",
            "Heading 2",
            "Heading 3",
            "Heading 4",
            "Heading 5",
            "Heading 6",
            "Title",
            "Subtitle",
        ]

    def _get_paragraph_text(self, para: Paragraph) -> str:
        return para.text.strip()

    def _get_paragraph_style(self, para: Paragraph) -> str:
        if para.style and para.style.name:
            return para.style.name
        return "Normal"

    def _is_heading(self, style_name: str) -> bool:
        return style_name in self.heading_styles or style_name.startswith("Heading")

    def _get_heading_level(self, style_name: str) -> int:
        if style_name == "Title":
            return 1
        if style_name == "Subtitle":
            return 2
        if style_name.startswith("Heading "):
            try:
                return int(style_name.split(" ")[-1])
            except ValueError:
                return 1
        return 0

    def _extract_table(self, table: Table) -> Dict[str, Any]:
        rows_data = []
        for row in table.rows:
            row_cells = []
            for cell in row.cells:
                row_cells.append(cell.text.strip())
            rows_data.append(row_cells)

        return {"type": "table", "data": rows_data}

    def _build_sections_recursive(
        self, elements: List, start_idx: int, current_level: int, parent_content: List
    ) -> tuple[List[SectionData], int]:
        sections = []
        i = start_idx

        while i < len(elements):
            element = elements[i]

            if isinstance(element, Paragraph):
                para_text = self._get_paragraph_text(element)
                style_name = self._get_paragraph_style(element)

                if self._is_heading(style_name):
                    if parent_content:
                        section = SectionData(
                            title=para_text,
                            level=current_level,
                            content=" ".join(parent_content),
                            children=[],
                        )
                        sections.append(section)
                        parent_content = []

                    heading_level = self._get_heading_level(style_name)
                    child_elements = elements[i + 1 :] if i + 1 < len(elements) else []
                    child_sections, next_idx = self._build_sections_recursive(
                        child_elements, 0, heading_level, []
                    )

                    section = SectionData(
                        title=para_text,
                        level=heading_level,
                        content="",
                        children=child_sections,
                    )
                    sections.append(section)
                    i += next_idx
                else:
                    if para_text:
                        parent_content.append(para_text)

            elif isinstance(element, Table):
                if parent_content:
                    table_data = self._extract_table(element)
                    parent_content.append(
                        f"[TABLE: {' | '.join([' | '.join(row) for row in table_data['data']])}]"
                    )

            i += 1

        if parent_content:
            section = SectionData(
                title="",
                level=current_level + 1,
                content=" ".join(parent_content),
                children=[],
            )
            sections.append(section)

        return sections, i

    def _extract_all_content(self, doc: Document) -> List:
        elements = []

        for element in doc.element.body:
            if isinstance(element, CT_P):
                para = Paragraph(element, doc)
                elements.append(para)
            elif isinstance(element, CT_Tbl):
                table = Table(element, doc)
                elements.append(table)

        return elements

    def parse(self, file_path: str) -> ParsedDocument:
        logger.info(f"Parsing DOCX file: {file_path}")

        doc = Document(file_path)
        doc_id = str(uuid.uuid4())

        elements = self._extract_all_content(doc)
        sections, _ = self._build_sections_recursive(elements, 0, 1, [])

        full_text_parts = []
        for section in sections:
            if section.content:
                full_text_parts.append(section.content)
            for child in section.children:
                if child.content:
                    full_text_parts.append(child.content)

        full_text = "\n\n".join(full_text_parts)

        metadata = {
            "paragraph_count": len(doc.paragraphs),
            "table_count": len(doc.tables),
            "page_count": getattr(doc, "page_count", "unknown"),
        }

        logger.info(f"Extracted {len(sections)} sections from DOCX")

        return ParsedDocument(
            doc_id=doc_id,
            filename=file_path.split("/")[-1].split("\\")[-1],
            file_type="docx",
            sections=sections,
            full_text=full_text,
            metadata=metadata,
        )


def parse_docx(file_path: str) -> ParsedDocument:
    parser = DocxParser()
    return parser.parse(file_path)
